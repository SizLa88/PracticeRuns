import openpyxl


class ExcelReader:

    FILE_PATH = "UIBank/TestData/RegistrationData.xlsx"

    @staticmethod
    def get_cell_data(row_num, col_num):

        try:
            workbook = openpyxl.load_workbook(
                ExcelReader.FILE_PATH
            )

            sheet = workbook.active

            cell_value = sheet.cell(
                row=row_num + 1,
                column=col_num + 1
            ).value

            return str(cell_value).strip() \
                if cell_value is not None else ""

        except Exception as e:

            print(f"Excel Read Error: {e}")

            return ""

    @staticmethod
    def get_row_count():

        try:
            workbook = openpyxl.load_workbook(
                ExcelReader.FILE_PATH
            )

            sheet = workbook.active

            return sheet.max_row - 1

        except Exception as e:

            print(f"Excel Read Error: {e}")

            return 0
